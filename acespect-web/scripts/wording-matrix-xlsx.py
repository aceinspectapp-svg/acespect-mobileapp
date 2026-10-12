#!/usr/bin/env python3
"""Builds the review spreadsheet from the JSON that scripts/wording-matrix.mjs writes.

    python3 scripts/wording-matrix-xlsx.py wording-matrix-residential_house.json [out.xlsx] [--refs a.txt,b.txt]

--refs: plain-text copies of the reference documents (Word templates / example reports). Each generated sentence is then checked
against the closest sentence in them: "Matches", "Close" (shown with the reference sentence) or nothing to compare with.

Needs openpyxl (pip install openpyxl).
"""
import json
import math
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "183966"
RED = "CE2328"
import difflib
import re

argv = [a for a in sys.argv[1:] if not a.startswith("--refs")]
ref_arg = next((a.split("=", 1)[1] if "=" in a else sys.argv[sys.argv.index(a) + 1] for a in sys.argv[1:] if a.startswith("--refs")), "")
if "--refs" in sys.argv:  # "--refs a,b" (two words) leaves the list in argv; drop it
    argv = [a for a in argv if a != ref_arg]
src = argv[0]
out = argv[1] if len(argv) > 1 else src.replace(".json", ".xlsx")
d = json.load(open(src))

# ---- reference sentences ------------------------------------------------------
def ref_sentences(paths):
    found = []
    for path in [p for p in ref_arg.split(",") if p]:
        text = open(path, errors="ignore").read()
        text = re.sub(r"[ \t]+", " ", text)
        for line in text.split("\n"):
            for sent in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", line.strip()):
                words = norm_words(sent)
                if len(words) >= 3:
                    found.append((sent.strip(), words, path.split("/")[-1]))
    return found


def norm_words(sent):
    return re.findall(r"[a-z0-9]+", sent.lower().replace("&", " and "))


REFS = ref_sentences(ref_arg.split(",")) if ref_arg else []
REF_INDEX = {}
for i, (_, words, _) in enumerate(REFS):
    for w in set(words):
        REF_INDEX.setdefault(w, []).append(i)
SKIP_LINE = re.compile(r"^\[(Condition tag|Heading|Scope page)")


def ref_check(sentence):
    """('Matches'|'Close'|'', reference sentence, source) for the closest reference sentence to each generated sentence."""
    if not REFS:
        return "", "", ""
    results = []
    for line in sentence.split("\n"):
        line = re.sub(r"^\[Scope page\]\s*", "", line)
        if not line or line.startswith("[") or line.startswith("(") or line.isupper():
            continue
        for sent in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", line):
            gen = norm_words(sent)
            if len(gen) < 3:
                continue
            votes = {}
            for w in set(gen):
                for i in REF_INDEX.get(w, []):
                    votes[i] = votes.get(i, 0) + 1
            best = (0.0, None, 0.0)
            for i, _ in sorted(votes.items(), key=lambda kv: -kv[1])[:12]:
                ref_words = REFS[i][1]
                m = difflib.SequenceMatcher(None, gen, ref_words, autojunk=False)
                matched = sum(b.size for b in m.get_matching_blocks())
                cover_ref = matched / len(ref_words)
                cover_gen = matched / len(gen)
                score = min(cover_ref, cover_gen) if len(ref_words) > 5 else m.ratio()
                score = max(score, m.ratio())
                if cover_ref >= 0.9 and cover_gen >= 0.55 and len(ref_words) >= 4:  # the reference has blanks the answer fills in
                    score = max(score, 0.95)
                if score > best[0]:
                    best = (score, i, m.ratio())
            if best[1] is None:
                results.append(("", "", ""))
            elif best[0] >= 0.95:
                results.append(("Matches", REFS[best[1]][0], REFS[best[1]][2]))
            elif best[0] >= 0.6:
                results.append(("Close", REFS[best[1]][0], REFS[best[1]][2]))
            else:
                results.append(("", "", ""))
    if not results:
        return "", "", ""
    rank = {"": 0, "Close": 1, "Matches": 2}
    status = "Matches" if all(r[0] == "Matches" for r in results) else ("Close" if any(r[0] for r in results) else "")
    if status == "Matches":
        return status, "", results[0][2]
    shown = [f"{r[1]}" for r in results if r[0] == "Close"]
    return status, "\n".join(dict.fromkeys(shown)), next((r[2] for r in results if r[0]), "")

TITLES = {
    "residential_house": "Dilapidation – Residential House",
    "commercial_properties": "Dilapidation – Commercial Properties",
    "apartment": "Dilapidation – Apartment",
    "public_assets": "Dilapidation – Public Assets",
}
report = TITLES.get(d["profile"].split(" / ")[-1], d["profile"])

thin = Side(style="thin", color="D9DEE6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FILLS = [PatternFill("solid", fgColor=c) for c in ("EAEFF6", "FFFFFF")]
WRAP = Alignment(wrap_text=True, vertical="top")


def est_height(values_widths):
    lines = 1
    for text, width in values_widths:
        n = 0
        for part in str(text).split("\n"):
            n += max(1, math.ceil(len(part) / max(width * 1.05, 1)))
        lines = max(lines, n)
    return min(15 * lines + 4, 400)


def sheet(wb, title, headers, widths, rows, review=True, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    cols = headers + (["Result", "Reviewer notes"] if review else [])
    ws_widths = widths + ([12, 42] if review else [])
    for i, (h, w) in enumerate(zip(cols, ws_widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    band, last_section = 0, None
    for r, row in enumerate(rows, 2):
        section = row[1] if len(row) > 1 else None
        if section != last_section:
            band, last_section = 1 - band, section
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            c.border = BORDER
            c.fill = SECTION_FILLS[band]
        if review:
            for i in (len(headers) + 1, len(headers) + 2):
                c = ws.cell(row=r, column=i)
                c.alignment = WRAP
                c.border = BORDER
        ws.row_dimensions[r].height = est_height(zip(row, ws_widths))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"
    if review and rows:
        col = get_column_letter(len(headers) + 1)
        rng = f"{col}2:{col}{len(rows) + 1}"
        dv = DataValidation(type="list", formula1='"Pass,Fix,Question"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng)
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", bgColor="C6EFCE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fix"'], fill=PatternFill("solid", bgColor="FFC7CE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Question"'], fill=PatternFill("solid", bgColor="FFEB9C")))
    return ws


wb = Workbook()

# ---- Read me ---------------------------------------------------------------
ws = wb.active
ws.title = "Read me"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110
lines = [
    ("title", f"{report} – wording matrix"),
    ("", f"Generated {d['generatedAt'][:10]} from the published form and the report's own wording code, so this is exactly what the report prints."),
    ("", ""),
    ("head", "How to review"),
    ("Answers", "Every answer an inspector can choose, and the sentence(s) the report prints for it. Read each row: is the sentence right, natural and complete? Set Result to Pass, Fix or Question and put the correction in Reviewer notes."),
    ("Defects", "Every defect type and sub-type, and how its sentence reads. Defect wording is the same in every section, so each is listed once."),
    ("Section samples", "A whole section filled in, so you can check that the sentences read well together (the Answers sheet shows single answers). Each sample lists the answers used."),
    ("Not in the wording", "Form fields that never change a sentence (photos, and any field the wording does not use). Check each is meant to be left out."),
    ("", ""),
    ("head", "How to read a row"),
    ("Answer", "The option picked on the form (or text typed: “Typed example” stands for whatever the inspector types)."),
    ("Sentence", "What the report adds for that answer. “(no change to the wording)” means the answer prints nothing by itself (for example “No”)."),
    ("Applies to", "Rooms, sides or areas that word the answer the same way are merged into one row; the sentence shown is the first, and the others use their own name in the same place. Blank means the section has one place only."),
    ("Shown", "“With the rest of the section answered” means the sentence only exists alongside other answers (the Description's one flowing paragraph), so the rest of the section was filled in to show it."),
    ("Reference document", "Each sentence is compared with the closest sentence in the reference documents supplied: “Matches” (the same wording, filling in the blanks), “Close” (similar; the reference sentence is shown so the difference is visible) or “No equivalent” (the references have no sentence like it, so it needs a judgement on its own). The references are templates with blanks, so “No equivalent” is common and is not a fault."),
    ("Check", "An automatic flag for something that reads wrongly (repeated word, “&”, a sentence ending on “to”, stray punctuation). A blank does not mean the row is right; it only means nothing mechanical was found."),
    ("[Scope page]", "A sentence printed under “Scope of Inspection and Comments” on the Description page, not in the property paragraph."),
    ("[Heading: …]", "A room or area heading rather than a sentence."),
    ("", ""),
    ("head", "What this does not cover"),
    ("", "Fixed text that never changes with an answer (cover, Scope appendix, Conditions of the report), photo layout and page breaks, and combinations of several answers beyond the section samples. Those are checked in the scenario and PDF review steps."),
    ("", ""),
    ("head", "Sizes"),
]
counts = Counter(r["section"] for r in d["matrix"])
for sec, n in counts.items():
    lines.append((sec, f"{n} answer rows"))
lines.append(("Defects", f"{len(d['defects'])} rows"))
lines.append(("Section samples", f"{len(d['samples'])} samples"))
lines.append(("Flagged rows", f"{sum(1 for r in d['matrix'] if r['flags'])} answer rows and {sum(1 for r in d['defects'] if r['flags'])} defect rows have an automatic flag"))
for i, (a, b) in enumerate(lines, 1):
    ca, cb = ws.cell(row=i, column=1, value=a if a not in ("title", "head") else b if a == "head" else b), ws.cell(row=i, column=2, value=b if a not in ("title", "head") else None)
    if a == "title":
        ca.font = Font(bold=True, size=16, color=NAVY)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    elif a == "head":
        ca.font = Font(bold=True, size=12, color="FFFFFF")
        ca.fill = HEAD_FILL
        ws.cell(row=i, column=2).fill = HEAD_FILL
    else:
        ca.font = Font(bold=True)
    ca.alignment = WRAP
    cb.alignment = WRAP
    if b and a != "title":
        ws.row_dimensions[i].height = est_height([(b, 110)])

# ---- Answers ----------------------------------------------------------------
def refcols(sentence):
    st, shown, src = ref_check(sentence)
    return (st or ("No equivalent" if REFS else "")), shown

ans_rows = []
for i, r in enumerate(d["matrix"], 1):
    st, ref = refcols(r["sentence"])
    ans_rows.append([i, r["section"], r["field"], r["answer"], r["sentence"], r["appliesTo"], "with the rest of the section answered" if r["context"] else "", r["summary"], r["flags"], st, ref])
heads = ["#", "Section", "Field (as the inspector sees it)", "Answer", "Sentence the report prints", "Applies to", "Shown", "Condition Summary line", "Check"]
widths = [5, 20, 34, 30, 72, 22, 16, 26, 18]
if REFS:
    heads += ["Reference document", "Closest reference sentence"]
    widths += [16, 56]
else:
    ans_rows = [r[:9] for r in ans_rows]
ws = sheet(wb, "Answers", heads, widths, ans_rows)
rows = ans_rows
for r in range(2, len(rows) + 2):
    if ws.cell(row=r, column=9).value:
        ws.cell(row=r, column=9).font = Font(bold=True, color=RED)
    if REFS:
        st = ws.cell(row=r, column=10).value
        ws.cell(row=r, column=10).fill = PatternFill("solid", fgColor={"Matches": "C6EFCE", "Close": "FFEB9C"}.get(st, "F2F2F2"))
    if str(ws.cell(row=r, column=5).value).startswith("(no change"):
        ws.cell(row=r, column=5).font = Font(italic=True, color="8A94A6")

# ---- Defects ------------------------------------------------------------------
rows = []
for i, r in enumerate(d["defects"], 1):
    st, ref = refcols(r["sentence"])
    rows.append([i, r["appliesTo"], r["situation"], r["answer"], r["sentence"], r["flags"]] + ([st, ref] if REFS else []))
ws = sheet(wb, "Defects", ["#", "Applies to", "Situation", "Defect / detail", "Sentence the report prints", "Check"] + (["Reference document", "Closest reference sentence"] if REFS else []), [5, 18, 30, 40, 72, 18] + ([16, 56] if REFS else []), rows)
for r in range(2, len(rows) + 2):
    if ws.cell(row=r, column=6).value:
        ws.cell(row=r, column=6).font = Font(bold=True, color=RED)
    if REFS:
        st = ws.cell(row=r, column=7).value
        ws.cell(row=r, column=7).fill = PatternFill("solid", fgColor={"Matches": "C6EFCE", "Close": "FFEB9C"}.get(st, "F2F2F2"))

# ---- Section samples ----------------------------------------------------------------
rows = [[i, r["section"], r["sample"], r["answers"], r["text"], r["summary"]] for i, r in enumerate(d["samples"], 1)]
sheet(wb, "Section samples", ["#", "Section", "Sample", "Answers used", "What the report prints", "Condition Summary lines"], [5, 20, 26, 48, 78, 30], rows)

# ---- Not in the wording ---------------------------------------------------------------
rows = [[i, r["section"], r["field"], "No answer changes any sentence – check that is intended"] for i, r in enumerate(d["unused"], 1)]
rows += [[len(d["unused"]) + i, r["section"], r["field"], "Photos (printed as photographs, not as a sentence)"] for i, r in enumerate(d["nonText"], 1)]
sheet(wb, "Not in the wording", ["#", "Section", "Field", "Why"], [5, 20, 60, 60], rows)

wb.save(out)
print(f"wrote {out}")
